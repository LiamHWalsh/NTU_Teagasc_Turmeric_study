import csv
import os
import re
from collections import defaultdict

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


BASE_DIR = os.path.join("results", "Processed", "05_statistical_analysis", "gps_injury_associations")
OUT_XLSX = os.path.join(BASE_DIR, "gps_microbiome_alignment_bundle_n11.xlsx")

GPS_COHORT = os.path.join(BASE_DIR, "gps_injury_analysis_cohort_n11.csv")
ALPHA_DELTAS = os.path.join(BASE_DIR, "participant_alpha_deltas_n11.csv")
SIGNAL_DELTAS = os.path.join(BASE_DIR, "participant_signal_deltas_n11.csv")

ALPHA_FILE = os.path.join("results", "Reports", "alpha_diversity.csv")
SPECIES_FILE = os.path.join("results", "Reports", "species_profile_filtered.csv")
PATHWAY_FILE = os.path.join("results", "Reports", "HUMAnN_merged_pathabundance_cpm.tsv")

STRAIN_EFF_FILE = os.path.join("results", "Processed", "05_statistical_analysis", "honest_small_n_strain", "strain_effect_sizes.csv")
GENUS_EFF_FILE = os.path.join("results", "Processed", "05_statistical_analysis", "honest_small_n_strain", "genus_effect_sizes.csv")
PATHWAY_EFF_FILE = os.path.join("results", "Processed", "05_statistical_analysis", "honest_small_n_strain", "pathway_effect_sizes.csv")
GPS_SOURCE_XLSX = os.getenv(
    "GPS_INJURY_XLSX",
    r"C:\Users\LiamWalsh\Downloads\OneDrive_1_3-1-2026\GPS and injury data_23_10.xlsx",
)

def read_csv_rows(path, delimiter=","):
    with open(path, newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle, delimiter=delimiter))


def read_header(path, delimiter=","):
    with open(path, newline="", encoding="utf-8-sig") as handle:
        return next(csv.reader(handle, delimiter=delimiter))


def normalize_pathway_sample_id(value):
    value = re.sub(r"_L001.*$", "", value)
    value = re.sub(r"_Abundance-CPM$", "", value)
    return value


def parse_sample_id(sample_id):
    match = re.match(r"^Nottingham_(\d+)([A-Za-z])_S\d+$", sample_id)
    if not match:
        return None
    block_letter = match.group(2).upper()
    timepoint = {
        "A": "Baseline",
        "B": "2_weeks",
        "C": "6_months",
    }.get(block_letter)
    return {
        "sample_id": sample_id,
        "participant_id": str(int(match.group(1))),
        "block_letter": block_letter,
        "timepoint": timepoint,
    }


def sort_effect_rows(rows):
    def key_fn(row):
        try:
            rank = float(row.get("rank_consensus", "inf"))
        except ValueError:
            rank = float("inf")
        try:
            effect = abs(float(row.get("cohens_d", "0")))
        except ValueError:
            effect = 0.0
        return (rank, -effect)

    return sorted(rows, key=key_fn)


def write_sheet(workbook, title, rows, header=None):
    ws = workbook.create_sheet(title=title)
    if header is None:
        header = list(rows[0].keys()) if rows else []
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([row.get(col, "") for col in header])

    for idx, col in enumerate(header, start=1):
        max_len = len(str(col))
        for cell in ws[get_column_letter(idx)]:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[get_column_letter(idx)].width = min(max_len + 2, 50)
    ws.freeze_panes = "A2"
    return ws


def long_to_wide(rows):
    outcomes = sorted({row["outcome"] for row in rows})
    by_participant = defaultdict(dict)
    for row in rows:
        pid = row["participant_id"]
        by_participant[pid][row["outcome"]] = row["delta"]

    wide_rows = []
    for participant_id in sorted(by_participant, key=lambda x: int(x)):
        row = {"participant_id": participant_id}
        for outcome in outcomes:
            row[outcome] = by_participant[participant_id].get(outcome, "")
        wide_rows.append(row)
    return wide_rows, ["participant_id"] + outcomes


def build_sample_map(gps_rows):
    analysis_ids = {row["participant_id"] for row in gps_rows}
    gps_lookup = {row["participant_id"]: row for row in gps_rows}

    alpha_header = read_header(ALPHA_FILE)
    alpha_rows = [parse_sample_id(row["sample_id"]) for row in read_csv_rows(ALPHA_FILE)]
    alpha_rows = [row for row in alpha_rows if row and row["participant_id"] in analysis_ids]

    species_samples = set(read_header(SPECIES_FILE)[1:])
    pathway_samples = {normalize_pathway_sample_id(x) for x in read_header(PATHWAY_FILE, delimiter="\t")[1:]}

    sample_rows = []
    wide_map = defaultdict(dict)
    for row in sorted(alpha_rows, key=lambda r: (int(r["participant_id"]), r["timepoint"])):
        pid = row["participant_id"]
        wide_map[pid][row["timepoint"]] = row["sample_id"]
        sample_rows.append({
            "participant_id": pid,
            "status": gps_lookup[pid]["status"],
            "turmeric_group": gps_lookup[pid]["turmeric_group"],
            "sample_id": row["sample_id"],
            "timepoint": row["timepoint"],
            "in_alpha": "TRUE",
            "in_species": "TRUE" if row["sample_id"] in species_samples else "FALSE",
            "in_pathway": "TRUE" if row["sample_id"] in pathway_samples else "FALSE",
        })

    wide_rows = []
    for pid in sorted(wide_map, key=lambda x: int(x)):
        wide_rows.append({
            "participant_id": pid,
            "status": gps_lookup[pid]["status"],
            "turmeric_group": gps_lookup[pid]["turmeric_group"],
            "Baseline_sample": wide_map[pid].get("Baseline", ""),
            "Week2_sample": wide_map[pid].get("2_weeks", ""),
            "Month6_sample": wide_map[pid].get("6_months", ""),
        })

    return sample_rows, wide_rows


def build_signal_lookup():
    lookup_rows = []

    strain_rows = sort_effect_rows(read_csv_rows(STRAIN_EFF_FILE))
    for row in strain_rows:
        lookup_rows.append({
            "level": "strain",
            "outcome": "strain_" + row["feature_short"],
            "feature_id": row["feature"],
            "feature_short": row["feature_short"],
            "rank_consensus": row["rank_consensus"],
            "cohens_d": row["cohens_d"],
        })

    genus_rows = sort_effect_rows(read_csv_rows(GENUS_EFF_FILE))
    for row in genus_rows:
        lookup_rows.append({
            "level": "genus",
            "outcome": "genus_" + row["feature_short"],
            "feature_id": row["feature"],
            "feature_short": row["feature_short"],
            "rank_consensus": row["rank_consensus"],
            "cohens_d": row["cohens_d"],
        })

    pathway_rows = sort_effect_rows(read_csv_rows(PATHWAY_EFF_FILE))
    for row in pathway_rows:
        lookup_rows.append({
            "level": "pathway",
            "outcome": "pathway_" + row["feature_short"],
            "feature_id": row["feature"],
            "feature_short": row["feature_short"],
            "rank_consensus": row["rank_consensus"],
            "cohens_d": row["cohens_d"],
        })

    return lookup_rows


def build_combined_rows(gps_rows, alpha_wide, signal_wide):
    alpha_lookup = {row["participant_id"]: row for row in alpha_wide}
    signal_lookup = {row["participant_id"]: row for row in signal_wide}

    rows = []
    for gps_row in sorted(gps_rows, key=lambda x: int(x["participant_id"])):
        pid = gps_row["participant_id"]
        merged = dict(gps_row)
        for key, value in alpha_lookup.get(pid, {}).items():
            if key != "participant_id":
                merged[key] = value
        for key, value in signal_lookup.get(pid, {}).items():
            if key != "participant_id":
                merged[key] = value
        rows.append(merged)
    return rows


def build_gps_source_header_rows():
    workbook = load_workbook(GPS_SOURCE_XLSX, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    header_rows = []
    for row_idx, row in enumerate(worksheet.iter_rows(min_row=1, max_row=2, values_only=True), start=1):
        for col_idx, value in enumerate(row, start=1):
            header_rows.append({
                "header_row": row_idx,
                "excel_column": get_column_letter(col_idx),
                "value": "" if value is None else value,
            })
    workbook.close()
    return header_rows


def build_gps_column_dictionary_rows():
    mappings = [
        ("participant_id", "A", "participant identifier", "", "direct import"),
        ("pre_training_total", "B", "Pre turmeric (14-week pre-season up to Visit 2)", "Training", "direct import"),
        ("pre_training_min_week", "C", "Pre turmeric (14-week pre-season up to Visit 2)", "Training minutes per week", "direct import"),
        ("pre_match_total", "D", "Pre turmeric (14-week pre-season up to Visit 2)", "Match", "direct import"),
        ("pre_match_min_week", "E", "Pre turmeric (14-week pre-season up to Visit 2)", "Match minutes per week", "direct import"),
        ("post_training_total", "G", "Post turmeric period (visit 3) (24-week season up to visit 3)", "Training Total", "direct import"),
        ("post_training_min_week", "H", "Post turmeric period (visit 3) (24-week season up to visit 3)", "Training minutes per week", "direct import"),
        ("post_match_total", "I", "Post turmeric period (visit 3) (24-week season up to visit 3)", "Match Total", "direct import"),
        ("post_match_min_week", "J", "Post turmeric period (visit 3) (24-week season up to visit 3)", "Match minutes per week", "direct import"),
        ("injury_days_pre", "Y", "Days absent injury", "pre turmeric", "direct import"),
        ("injury_pct_pre", "Z", "Days absent injury", "% days absent", "direct import"),
        ("injury_days_post", "AA", "Days absent injury", "post turmeric", "direct import"),
        ("injury_pct_post", "AB", "Days absent injury", "% days absent", "direct import"),
        ("illness_days_pre", "AC", "Days absent illness", "pre turmeric", "direct import"),
        ("illness_pct_pre", "AD", "Days absent illness", "% days absent", "direct import"),
        ("illness_days_post", "AE", "Days absent illness", "post turmeric", "direct import"),
        ("illness_pct_post", "AF", "Days absent illness", "% days absent", "direct import"),
        ("delta_training_total", "", "", "", "derived as post_training_total - pre_training_total"),
        ("delta_training_min_week", "", "", "", "derived as post_training_min_week - pre_training_min_week"),
        ("delta_match_total", "", "", "", "derived as post_match_total - pre_match_total"),
        ("delta_match_min_week", "", "", "", "derived as post_match_min_week - pre_match_min_week"),
        ("injury_days_total", "", "", "", "derived as injury_days_pre + injury_days_post"),
        ("illness_days_total", "", "", "", "derived as illness_days_pre + illness_days_post"),
    ]
    return [
        {
            "exported_column": exported_column,
            "source_excel_column": source_excel_column,
            "source_period_header": source_period_header,
            "source_subheader": source_subheader,
            "derivation": derivation,
        }
        for exported_column, source_excel_column, source_period_header, source_subheader, derivation in mappings
    ]


def main():
    gps_rows = read_csv_rows(GPS_COHORT)
    alpha_rows = read_csv_rows(ALPHA_DELTAS)
    signal_rows = read_csv_rows(SIGNAL_DELTAS)

    alpha_wide, alpha_header = long_to_wide(alpha_rows)
    signal_wide, signal_header = long_to_wide(signal_rows)
    sample_rows, sample_wide_rows = build_sample_map(gps_rows)
    signal_lookup_rows = build_signal_lookup()
    combined_rows = build_combined_rows(gps_rows, alpha_wide, signal_wide)
    gps_source_header_rows = build_gps_source_header_rows()
    gps_column_dictionary_rows = build_gps_column_dictionary_rows()

    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    readme_rows = [
        {
            "item": "cohort",
            "details": "Participant-level GPS/injury and microbiome alignment bundle for the n=11 treatment/control cohort used in the GPS association analysis.",
        },
        {
            "item": "delta_definition_alpha_signals",
            "details": "For GPS/injury associations, pre_value = mean(Baseline, 2_weeks), post_value = Month6, delta = post_value - pre_value.",
        },
        {
            "item": "sample_map",
            "details": "Sample IDs are parsed from alpha_diversity.csv and cross-checked against species_profile_filtered.csv and HUMAnN_merged_pathabundance_cpm.tsv headers.",
        },
        {
            "item": "combined_sheet",
            "details": "One row per participant with GPS/injury variables, alpha deltas, and selected microbiome signal deltas used in the association analysis.",
        },
        {
            "item": "gps_period_mapping",
            "details": "GPS variables use flattened names in the export: pre_* columns map to 'Pre turmeric (14-week pre-season up to Visit 2)' and post_* columns map to 'Post turmeric period (visit 3) (24-week season up to visit 3)'. See sheets gps_source_headers and gps_column_dictionary.",
        },
    ]

    write_sheet(workbook, "README", readme_rows, header=["item", "details"])
    write_sheet(workbook, "gps_source_headers", gps_source_header_rows)
    write_sheet(workbook, "gps_column_dictionary", gps_column_dictionary_rows)
    write_sheet(workbook, "gps_analysis_n11", gps_rows)
    write_sheet(workbook, "sample_map_n11_long", sample_rows)
    write_sheet(workbook, "sample_map_n11_wide", sample_wide_rows)
    write_sheet(workbook, "alpha_deltas_long", alpha_rows)
    write_sheet(workbook, "alpha_deltas_wide", alpha_wide, header=alpha_header)
    write_sheet(workbook, "signal_deltas_long", signal_rows)
    write_sheet(workbook, "signal_deltas_wide", signal_wide, header=signal_header)
    write_sheet(workbook, "signal_feature_lookup", signal_lookup_rows)
    write_sheet(workbook, "combined_n11", combined_rows)

    workbook.save(OUT_XLSX)
    print(OUT_XLSX)


if __name__ == "__main__":
    main()
